import json
from pathlib import Path

from openpyxl import load_workbook

from clean_databricks.configuration import AnalysisSettings
from clean_databricks.dependence import _parse_references, run_dependence_analysis


def test_parser_resolves_relations_and_excludes_ctes() -> None:
    references = _parse_references(
        """
        CREATE VIEW current_view AS
        WITH recent AS (SELECT * FROM source_table)
        SELECT * FROM recent JOIN other_schema.lookup USING (id)
        """,
        "main",
        "analytics",
    )

    assert sorted(references, key=lambda item: item["normalized_name"]) == [
        {
            "original_text": "source_table",
            "normalized_name": "main.analytics.source_table",
            "category": "absent_from_inventory",
        },
        {
            "original_text": "other_schema.lookup",
            "normalized_name": "main.other_schema.lookup",
            "category": "outside_scope",
        },
    ]


def test_dependence_analysis_writes_human_and_machine_artifacts(
    tmp_path: Path,
) -> None:
    inventory = tmp_path / "inventory"
    raw = inventory / "raw"
    raw.mkdir(parents=True)
    objects = {
        "source-id": {
            "brick_id": "source-id",
            "full_name": "main.analytics.source_table",
            "family": "table",
            "kind": "managed_table",
            "status": "complete",
            "file": "raw/source_table.json",
        },
        "view-id": {
            "brick_id": "view-id",
            "full_name": "main.analytics.report_view",
            "family": "view",
            "kind": "view",
            "status": "incomplete",
            "file": "raw/report_view.json",
        },
    }
    (inventory / "manifest.json").write_text(
        json.dumps(
            {
                "format_version": 1,
                "scope": {"catalog": "main", "schema": "analytics"},
                "objects": objects,
                "last_run_status": "partial",
            }
        ),
        encoding="utf-8",
    )
    for filename, sql in (
        ("source_table.json", "CREATE TABLE source_table (id INT)"),
        (
            "report_view.json",
            "CREATE VIEW report_view AS SELECT * FROM source_table",
        ),
    ):
        (raw / filename).write_text(
            json.dumps(
                {
                    "collection": {
                        "show_create": {
                            "status": "collected",
                            "columns": ["createtab_stmt"],
                            "rows": [[sql]],
                        }
                    }
                }
            ),
            encoding="utf-8",
        )
    settings = AnalysisSettings(
        "main", "analytics", inventory, tmp_path / "analysis"
    )

    output = run_dependence_analysis(settings)

    result = json.loads((output / "dependence.json").read_text(encoding="utf-8"))
    assert len(result["edges"]) == 1
    assert result["edges"][0]["source_id"] == "view-id"
    assert result["edges"][0]["target_id"] == "source-id"
    for relative in (
        "summary.md",
        "dependence.xlsx",
        "nodes.csv",
        "edges.csv",
        "figures/overview.svg",
        "manifest.json",
    ):
        assert (output / relative).is_file()
    workbook = load_workbook(output / "dependence.xlsx")
    sheet = workbook["BrickObjects"]
    headers = [cell.value for cell in sheet[1]]
    raw_column = headers.index("raw_record") + 1
    assert sheet.cell(row=2, column=raw_column).hyperlink is not None
