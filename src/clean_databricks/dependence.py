"""Offline BrickObject dependence analysis and artifact generation."""

from __future__ import annotations

import csv
import json
import os
import shutil
import uuid
from collections import defaultdict
from collections.abc import Callable, Iterable, Mapping
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import networkx as nx
import sqlglot
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlglot import expressions as exp

from .configuration import AnalysisSettings

FORMAT_VERSION = 1


def run_dependence_analysis(
    settings: AnalysisSettings,
    progress: Callable[[str], None] | None = None,
) -> Path:
    """Analyze collected definitions and transactionally publish artifacts."""
    emit = progress or (lambda _message: None)
    emit("Loading Inventory...")
    _assert_inventory_idle(settings.inventory_dir)
    manifest_path = settings.inventory_dir / "manifest.json"
    if not manifest_path.is_file():
        raise ValueError(f"Inventory Manifest not found: {manifest_path}")
    manifest = _read_object(manifest_path, "Inventory Manifest")
    scope = manifest.get("scope")
    expected_scope = {"catalog": settings.catalog, "schema": settings.schema}
    if scope != expected_scope:
        raise ValueError(
            "Inventory Manifest scope does not match configured catalog and schema."
        )
    entries = manifest.get("objects")
    if not isinstance(entries, dict) or not entries:
        raise ValueError("Inventory Manifest contains no BrickObjects to analyze.")

    nodes: dict[str, dict[str, Any]] = {}
    diagnostics: list[dict[str, Any]] = []
    records: list[tuple[dict[str, Any], Path, dict[str, Any]]] = []
    for brick_id, value in sorted(entries.items()):
        if not isinstance(value, dict):
            diagnostics.append(_diagnostic(str(brick_id), "invalid_manifest_entry"))
            continue
        full_name = str(value.get("full_name", ""))
        node = {
            "id": str(brick_id),
            "full_name": full_name,
            "display_name": full_name,
            "kind": value.get("kind", "unknown"),
            "family": value.get("family", "unknown"),
            "resolution": "inventory",
            "raw_record": None,
        }
        nodes[str(brick_id)] = node
        relative = value.get("file")
        if not isinstance(relative, str):
            diagnostics.append(_diagnostic(str(brick_id), "missing_record_path"))
            continue
        raw_path = settings.inventory_dir / relative
        node["raw_record"] = _display_path(
            raw_path, settings.analysis_dir / "dependence"
        )
        try:
            record = _read_object(raw_path, "Raw Object Record")
        except (OSError, ValueError) as exc:
            diagnostics.append(
                _diagnostic(str(brick_id), "unreadable_record", str(exc))
            )
            continue
        records.append((node, raw_path, record))

    if not records:
        raise ValueError("Inventory contains no readable Raw Object Records.")

    emit("Parsing collected definitions...")
    known = {
        str(node["full_name"]).casefold(): node_id
        for node_id, node in nodes.items()
        if node["full_name"]
    }
    edge_evidence: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for node, raw_path, record in records:
        sql, operation = _definition_sql(record)
        if not sql:
            diagnostics.append(_diagnostic(node["id"], "definition_unavailable"))
            continue
        try:
            references = _parse_references(sql, settings.catalog, settings.schema)
        except sqlglot.errors.ParseError as exc:
            diagnostics.append(
                _diagnostic(node["id"], "parse_failure", str(exc).splitlines()[0])
            )
            continue
        for reference in references:
            normalized = reference["normalized_name"]
            target_id = known.get(normalized.casefold())
            if target_id is None:
                category = reference["category"]
                target_id = _unresolved_id(category, normalized)
                nodes.setdefault(
                    target_id,
                    {
                        "id": target_id,
                        "full_name": normalized,
                        "display_name": reference["original_text"],
                        "kind": "unresolved",
                        "family": "unresolved",
                        "resolution": category,
                        "raw_record": None,
                    },
                )
            edge_evidence[(str(node["id"]), target_id)].append(
                {
                    "source_record": _display_path(
                        raw_path, settings.analysis_dir / "dependence"
                    ),
                    "operation": operation,
                    "original_reference": reference["original_text"],
                    "normalized_target": normalized,
                }
            )

    emit("Resolving references and computing graph findings...")
    edges = [
        {
            "source_id": source,
            "target_id": target,
            "occurrence_count": len(evidence),
            "evidence": evidence,
        }
        for (source, target), evidence in sorted(edge_evidence.items())
    ]
    graph: nx.DiGraph[str] = nx.DiGraph()
    graph.add_nodes_from(nodes)
    for edge in edges:
        graph.add_edge(str(edge["source_id"]), str(edge["target_id"]))
    components = _components(graph)
    cycles = _cycles(graph)
    component_by_node = {
        node_id: component["id"]
        for component in components
        for node_id in component["node_ids"]
    }
    for node_id, node in nodes.items():
        node["direct_dependencies"] = graph.out_degree(node_id)
        node["direct_dependents"] = graph.in_degree(node_id)
        node["transitive_dependents"] = len(nx.ancestors(graph, node_id))
        node["component_id"] = component_by_node[node_id]

    result = {
        "format_version": FORMAT_VERSION,
        "generated_at": _now(),
        "scope": expected_scope,
        "inventory": {
            "path": _display_path(
                manifest_path, settings.analysis_dir / "dependence"
            ),
            "last_run_status": manifest.get("last_run_status"),
            "manifest_objects": len(entries),
            "readable_records": len(records),
        },
        "edge_direction": "source BrickObject depends on target BrickObject",
        "nodes": sorted(nodes.values(), key=lambda item: str(item["id"])),
        "edges": edges,
        "components": components,
        "cycles": cycles,
        "diagnostics": diagnostics,
    }

    emit("Writing analysis artifacts...")
    root = settings.analysis_dir
    root.mkdir(parents=True, exist_ok=True)
    staging = root / f".dependence-{uuid.uuid4().hex}.tmp"
    staging.mkdir()
    try:
        _write_json(staging / "dependence.json", result)
        _write_csv_artifacts(staging, result)
        _write_workbook(staging / "dependence.xlsx", result)
        _write_figures(staging, graph, nodes, components)
        (staging / "summary.md").write_text(
            _summary_markdown(result), encoding="utf-8"
        )
        artifacts = sorted(
            path.relative_to(staging).as_posix()
            for path in staging.rglob("*")
            if path.is_file()
        )
        _write_json(
            staging / "manifest.json",
            {
                "format_version": FORMAT_VERSION,
                "analysis": "dependence",
                "generated_at": result["generated_at"],
                "scope": expected_scope,
                "counts": {
                    "nodes": len(nodes),
                    "edges": len(edges),
                    "cycles": len(cycles),
                    "diagnostics": len(diagnostics),
                },
                "artifacts": [*artifacts, "manifest.json"],
            },
        )
        emit("Publishing dependence analysis...")
        target = root / "dependence"
        backup = root / f".dependence-{uuid.uuid4().hex}.bak"
        if target.exists():
            target.rename(backup)
        try:
            staging.rename(target)
        except Exception:
            if backup.exists():
                backup.rename(target)
            raise
        shutil.rmtree(backup, ignore_errors=True)
        emit(
            f"Analyzed {len(records)} records: {len(nodes)} nodes, "
            f"{len(edges)} edges, {len(cycles)} cycles, "
            f"{len(diagnostics)} diagnostics."
        )
        return target
    finally:
        shutil.rmtree(staging, ignore_errors=True)


def _definition_sql(record: Mapping[str, Any]) -> tuple[str | None, str | None]:
    collection = record.get("collection")
    if not isinstance(collection, dict):
        return None, None
    show_create = collection.get("show_create")
    if isinstance(show_create, dict) and show_create.get("status") == "collected":
        rows = show_create.get("rows")
        if isinstance(rows, list) and rows and isinstance(rows[0], list) and rows[0]:
            return str(rows[0][0]), "show_create"
    describe = collection.get("describe_extended")
    if isinstance(describe, dict) and describe.get("status") == "collected":
        for value in _walk_values(describe.get("rows")):
            if isinstance(value, str) and value.lstrip().upper().startswith("CREATE "):
                return value, "describe_extended"
    return None, None


def _walk_values(value: Any) -> Iterable[Any]:
    if isinstance(value, dict):
        for nested in value.values():
            yield from _walk_values(nested)
    elif isinstance(value, list):
        for nested in value:
            yield from _walk_values(nested)
    else:
        yield value


def _parse_references(sql: str, catalog: str, schema: str) -> list[dict[str, str]]:
    references: list[dict[str, str]] = []
    for statement in sqlglot.parse(sql, read="databricks"):
        search_root = (
            statement.expression
            if isinstance(statement, exp.Create)
            else statement
        )
        if search_root is None:
            continue
        ctes = {
            cte.alias_or_name.casefold()
            for cte in search_root.find_all(exp.CTE)
            if cte.alias_or_name
        }
        for table in search_root.find_all(exp.Table):
            name = table.name
            if not name or (not table.db and name.casefold() in ctes):
                continue
            original = table.sql(dialect="databricks")
            if table.db.casefold() in {"delta", "parquet", "csv", "json"} and (
                "/" in name or "\\" in name
            ):
                normalized = original
                category = "path_reference"
            else:
                target_catalog = table.catalog or catalog
                target_schema = table.db or schema
                normalized = f"{target_catalog}.{target_schema}.{name}"
                category = (
                    "absent_from_inventory"
                    if target_catalog.casefold() == catalog.casefold()
                    and target_schema.casefold() == schema.casefold()
                    else "outside_scope"
                )
            references.append(
                {
                    "original_text": original,
                    "normalized_name": normalized,
                    "category": category,
                }
            )
    return references


def _components(graph: nx.DiGraph[str]) -> list[dict[str, Any]]:
    groups = sorted(
        (sorted(group) for group in nx.weakly_connected_components(graph)),
        key=lambda group: (-len(group), group),
    )
    return [
        {
            "id": f"component-{index:04d}",
            "node_ids": group,
            "node_count": len(group),
            "edge_count": graph.subgraph(group).number_of_edges(),
        }
        for index, group in enumerate(groups, start=1)
    ]


def _cycles(graph: nx.DiGraph[str]) -> list[dict[str, Any]]:
    groups = []
    for group in nx.strongly_connected_components(graph):
        ordered = sorted(group)
        if len(ordered) > 1 or (ordered and graph.has_edge(ordered[0], ordered[0])):
            groups.append(ordered)
    return [
        {"id": f"cycle-{index:04d}", "node_ids": group}
        for index, group in enumerate(sorted(groups), start=1)
    ]


def _write_csv_artifacts(directory: Path, result: Mapping[str, Any]) -> None:
    node_fields = [
        "id", "full_name", "kind", "family", "resolution", "raw_record",
        "direct_dependencies", "direct_dependents", "transitive_dependents",
        "component_id",
    ]
    _write_csv(directory / "nodes.csv", node_fields, result["nodes"])
    edge_rows = [
        {key: edge[key] for key in ("source_id", "target_id", "occurrence_count")}
        for edge in result["edges"]
    ]
    _write_csv(
        directory / "edges.csv",
        ["source_id", "target_id", "occurrence_count"],
        edge_rows,
    )


def _write_csv(
    path: Path, fields: list[str], rows: Iterable[Mapping[str, Any]]
) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _write_workbook(path: Path, result: Mapping[str, Any]) -> None:
    workbook = Workbook()
    active = workbook.active
    if active is not None:
        workbook.remove(active)
    summary = [
        {"metric": "BrickObjects", "value": len(result["nodes"])},
        {"metric": "Dependency Edges", "value": len(result["edges"])},
        {"metric": "Cycles", "value": len(result["cycles"])},
        {"metric": "Diagnostics", "value": len(result["diagnostics"])},
    ]
    _sheet(workbook, "Summary", summary)
    _sheet(workbook, "BrickObjects", result["nodes"])
    _sheet(workbook, "Dependency Edges", result["edges"])
    _sheet(workbook, "Components", result["components"])
    _sheet(workbook, "Cycles", result["cycles"])
    unresolved = [node for node in result["nodes"] if node["resolution"] != "inventory"]
    _sheet(workbook, "Unresolved References", unresolved)
    _sheet(workbook, "Diagnostics", result["diagnostics"])
    workbook.save(path)


def _sheet(workbook: Workbook, title: str, rows: Iterable[Mapping[str, Any]]) -> None:
    values = list(rows)
    sheet = workbook.create_sheet(title)
    fields = list(values[0]) if values else ["message"]
    sheet.append(fields)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in values:
        sheet.append([_excel_value(row.get(field)) for field in fields])
        if title == "BrickObjects" and "raw_record" in fields:
            column = fields.index("raw_record") + 1
            cell = sheet.cell(row=sheet.max_row, column=column)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, field in enumerate(fields, start=1):
        maximum = max([len(field), *(len(str(row.get(field, ""))) for row in values)])
        sheet.column_dimensions[get_column_letter(index)].width = min(maximum + 2, 60)


def _excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _write_figures(
    directory: Path,
    graph: nx.DiGraph[str],
    nodes: Mapping[str, Mapping[str, Any]],
    components: list[dict[str, Any]],
) -> None:
    figures = directory / "figures"
    figures.mkdir()
    sizes = [component["node_count"] for component in components]
    figure, axis = plt.subplots(figsize=(10, 5))
    axis.bar(range(1, len(sizes) + 1), sizes)
    axis.set(title="Dependence components", xlabel="Component", ylabel="BrickObjects")
    figure.tight_layout()
    figure.savefig(figures / "overview.svg", format="svg")
    plt.close(figure)
    for component in components:
        if component["node_count"] < 2:
            continue
        subgraph = graph.subgraph(component["node_ids"])
        scale = max(8.0, min(30.0, component["node_count"] * 0.35))
        figure, axis = plt.subplots(figsize=(scale, scale))
        positions = nx.spring_layout(subgraph, seed=42)
        labels = {
            node_id: str(nodes[node_id]["display_name"]).split(".")[-1]
            for node_id in subgraph.nodes
        }
        colors = [
            "#d97706" if nodes[node_id]["resolution"] != "inventory" else "#2563eb"
            for node_id in subgraph.nodes
        ]
        nx.draw_networkx(
            subgraph,
            positions,
            labels=labels,
            node_color=colors,
            node_size=650,
            font_size=7,
            arrows=True,
            ax=axis,
        )
        axis.set_title(f"{component['id']} — arrow points to dependency")
        axis.axis("off")
        figure.tight_layout()
        figure.savefig(figures / f"{component['id']}.svg", format="svg")
        plt.close(figure)


def _summary_markdown(result: Mapping[str, Any]) -> str:
    inventory = result["inventory"]
    nodes = result["nodes"]
    unresolved = [node for node in nodes if node["resolution"] != "inventory"]
    top = sorted(
        (node for node in nodes if node["resolution"] == "inventory"),
        key=lambda node: (-node["transitive_dependents"], node["full_name"]),
    )[:20]
    lines = [
        "# Dependence Analysis",
        "",
        "> A graph edge points from a dependent BrickObject to its dependency.",
        "> No observed dependency does not mean that an object is unused.",
        "",
        "## Coverage",
        "",
        f"- Inventory run status: `{inventory['last_run_status']}`",
        f"- Manifest BrickObjects: {inventory['manifest_objects']}",
        f"- Readable Raw Object Records: {inventory['readable_records']}",
        f"- Graph nodes: {len(nodes)}",
        f"- Direct Dependency Edges: {len(result['edges'])}",
        f"- Diagnostics: {len(result['diagnostics'])}",
        "",
        "## Cycles",
        "",
        f"Found {len(result['cycles'])} cyclic component(s).",
        "",
        "## Unresolved references",
        "",
        f"Found {len(unresolved)} unresolved reference node(s).",
        "",
        "## Highest dependent impact",
        "",
        "| BrickObject | Direct dependents | Transitive dependents |",
        "|---|---:|---:|",
    ]
    lines.extend(
        f"| `{node['full_name']}` | {node['direct_dependents']} | "
        f"{node['transitive_dependents']} |"
        for node in top
    )
    lines.extend(
        [
            "",
            "## Artifacts",
            "",
            "- [Workbook](dependence.xlsx)",
            "- [Machine-readable graph](dependence.json)",
            "- [Nodes CSV](nodes.csv)",
            "- [Edges CSV](edges.csv)",
            "- [Figure overview](figures/overview.svg)",
        ]
    )
    return "\n".join(lines) + "\n"


def _assert_inventory_idle(inventory_dir: Path) -> None:
    lock_path = inventory_dir / "tmp" / "inventory.lock"
    if not lock_path.exists():
        return
    try:
        lock = json.loads(lock_path.read_text(encoding="utf-8"))
        pid = int(lock["pid"])
        os.kill(pid, 0)
    except (OSError, ValueError, KeyError, json.JSONDecodeError):
        return
    raise RuntimeError(
        f"Inventory is currently running (process {pid}); wait for it to finish."
    )


def _read_object(path: Path, label: str) -> dict[str, Any]:
    try:
        loaded = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Invalid {label} at {path}: {exc}") from exc
    if not isinstance(loaded, dict):
        raise ValueError(f"{label} at {path} must contain a JSON object.")
    return loaded


def _diagnostic(node_id: str, code: str, message: str = "") -> dict[str, str]:
    return {"node_id": node_id, "code": code, "message": message}


def _unresolved_id(category: str, normalized: str) -> str:
    import hashlib

    digest = hashlib.sha256(f"{category}:{normalized.casefold()}".encode()).hexdigest()
    return f"unresolved-{digest[:20]}"


def _display_path(path: Path, analysis_dir: Path) -> str:
    try:
        return os.path.relpath(path.resolve(), analysis_dir.resolve())
    except ValueError:
        return str(path.resolve())


def _write_json(path: Path, value: Any) -> None:
    path.write_text(
        json.dumps(value, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _now() -> str:
    return datetime.now(UTC).isoformat()
